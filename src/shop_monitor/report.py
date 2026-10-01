"""Отчёт: что изменилось и что сейчас на витрине — в консоль и в Excel."""

from __future__ import annotations

import os
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from shop_monitor.collect import Item
from shop_monitor.compare import APPEARED, DISAPPEARED, PRICE_DOWN, PRICE_UP, Change
from shop_monitor.profiles import Profile
from shop_monitor.storage import Run

SYMBOLS = {"RUB": "₽", "GBP": "£", "USD": "$", "EUR": "€", "KZT": "₸", "JPY": "¥"}
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5D50")
UP_FILL = PatternFill("solid", fgColor="F4CCCC")
DOWN_FILL = PatternFill("solid", fgColor="D9EAD3")
MONEY = "#,##0.00"
PERCENT = "+0.0%;−0.0%"
NBSP = " "


class ReportError(Exception):
    pass


def money(minor: int | None, currency: str) -> str:
    if minor is None:
        return "—"
    sign = "−" if minor < 0 else ""
    whole, fraction = divmod(abs(minor), 100)
    amount = f"{sign}{whole:,}".replace(",", NBSP) + f",{fraction:02d}"
    return f"{amount}{NBSP}{SYMBOLS.get(currency, currency)}".strip()


def stock(item: Item) -> str:
    if item.available is None:
        return "—"
    if not item.available:
        return "нет в наличии"
    return f"есть, {item.quantity} шт." if item.quantity else "есть"


def describe(change: Change) -> str:
    """Изменение словами: цена — стрелкой с процентом, наличие и состав витрины — текстом."""
    currency = change.item.currency
    if change.kind in (PRICE_UP, PRICE_DOWN):
        share = f" ({change.share:+.1%})".replace(".", ",").replace("-", "−") if change.share is not None else ""
        return f"{money(change.was, currency)} → {money(change.now, currency)}{share}"
    if change.kind == DISAPPEARED:
        return f"был по {money(change.was, currency)}"
    if change.kind == APPEARED:
        return f"{money(change.now, currency)}, {stock(change.item)}"
    return f"{stock(change.item)}, цена {money(change.now, currency)}"


def print_report(profile: Profile, items: Sequence[Item], changes: Sequence[Change], first_run: bool) -> None:
    print(f"{profile.title or profile.name}: {profile.url}")
    print(f"Собрано товаров: {len(items)}, страниц: {max((item.page for item in items), default=0)}")
    if first_run:
        print("Это первый прогон по этому профилю — сравнивать пока не с чем.")
    elif not changes:
        print("С прошлого прогона ничего не изменилось.")
    else:
        print(f"Изменений: {len(changes)}")
        for change in changes:
            print(f"  • {change.kind}: {change.item.name} — {describe(change)}")


def _header(sheet: Worksheet, titles: Sequence[str], widths: Sequence[float]) -> None:
    for column, (title, width) in enumerate(zip(titles, widths, strict=True), start=1):
        cell = sheet.cell(1, column, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"


def _changes_sheet(sheet: Worksheet, changes: Sequence[Change]) -> None:
    _header(
        sheet,
        ["Что произошло", "Товар", "Было", "Стало", "Разница", "Разница, %", "Наличие", "Ссылка"],
        [18, 42, 14, 14, 14, 12, 16, 50],
    )
    for row, change in enumerate(changes, start=2):
        values: list[str | int | float | Decimal | None] = [
            change.kind,
            change.item.name,
            None if change.was is None else Decimal(change.was) / 100,
            None if change.now is None else Decimal(change.now) / 100,
            None if change.difference is None else Decimal(change.difference) / 100,
            change.share,
            stock(change.item),
            change.item.url,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            if column in (3, 4, 5):
                cell.number_format = MONEY
            elif column == 6:
                cell.number_format = PERCENT
        if change.kind in (PRICE_UP, PRICE_DOWN):
            sheet.cell(row, 1).fill = UP_FILL if change.kind == PRICE_UP else DOWN_FILL


def _items_sheet(sheet: Worksheet, items: Sequence[Item], moves: dict[str, int]) -> None:
    _header(
        sheet,
        ["Позиция", "Сдвиг", "Товар", "Цена", "Валюта", "Наличие", "Остаток", "Рейтинг", "Страница", "Ссылка"],
        [9, 9, 42, 14, 9, 16, 10, 10, 10, 50],
    )
    for row, item in enumerate(items, start=2):
        values: list[str | int | float | Decimal | None] = [
            item.position,
            moves.get(item.key),
            item.name,
            None if item.price is None else Decimal(item.price) / 100,
            item.currency,
            stock(item),
            item.quantity,
            item.rating,
            item.page,
            item.url,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            if column == 4:
                cell.number_format = MONEY
    if items:
        sheet.auto_filter.ref = f"A1:J{len(items) + 1}"


def _runs_sheet(sheet: Worksheet, runs: Sequence[Run]) -> None:
    _header(sheet, ["Прогон", "Профиль", "Начат", "Товаров", "Страниц", "Адрес"], [9, 16, 20, 10, 10, 60])
    for row, run in enumerate(runs, start=2):
        sheet.cell(row, 1, run.id)
        sheet.cell(row, 2, run.profile)
        sheet.cell(row, 3, run.started.astimezone().replace(tzinfo=None)).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(row, 4, run.items)
        sheet.cell(row, 5, run.pages)
        sheet.cell(row, 6, run.url)


def write_report(
    path: Path,
    items: Sequence[Item],
    changes: Sequence[Change],
    runs: Sequence[Run] = (),
    moves: dict[str, int] | None = None,
) -> None:
    workbook = Workbook()
    del workbook["Sheet"]
    _changes_sheet(workbook.create_sheet("Изменения"), changes)
    _items_sheet(workbook.create_sheet("Витрина"), items, moves or {})
    if runs:
        _runs_sheet(workbook.create_sheet("Прогоны"), runs)
    temporary = path.with_name(path.name + ".tmp")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(temporary)
        os.replace(temporary, path)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise ReportError(f"Не удалось сохранить {path.name}: файл открыт в Excel? Закройте его и повторите.") from exc
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise ReportError(f"Не удалось сохранить {path.name}: {exc}") from exc
