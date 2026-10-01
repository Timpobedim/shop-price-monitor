from __future__ import annotations

from pathlib import Path

import pytest
from conftest import Shop
from openpyxl import load_workbook

from shop_monitor import demo
from shop_monitor.cli import main
from shop_monitor.profiles import demo_profile, to_json
from shop_monitor.report import money

NBSP = "\u00a0"


def test_money_formatting() -> None:
    assert money(249_000, "RUB") == f"2{NBSP}490,00{NBSP}₽"
    assert money(4782, "GBP") == f"47,82{NBSP}£"
    assert money(-30_000, "RUB") == f"−300,00{NBSP}₽"
    assert money(None, "RUB") == "—"


@pytest.mark.browser
def test_demo_command_runs_twice_and_reports_changes(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["demo", "--out", str(tmp_path)]) == 0
    printed = capsys.readouterr().out
    assert "Прогон 1" in printed
    assert "подорожал: Чайник электрический 1,7 л" in printed
    assert "исчез: Термос 1 л" in printed
    workbook = load_workbook(tmp_path / "отчёт.xlsx")
    assert workbook.sheetnames == ["Изменения", "Витрина", "Прогоны"]
    changes = [row[0] for row in workbook["Изменения"].iter_rows(min_row=2, values_only=True)]
    assert set(changes) == {"подорожал", "подешевел", "закончился", "снова в наличии", "появился", "исчез"}
    assert workbook["Витрина"].max_row == 13  # заголовок и 12 товаров
    assert workbook["Прогоны"].max_row == 3


@pytest.mark.browser
def test_run_with_own_profile_file(shop: Shop, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    profile = tmp_path / "профиль.json"
    profile.write_text(to_json(demo_profile(shop.url)), encoding="utf-8")
    database, report = tmp_path / "история.sqlite", tmp_path / "отчёт.xlsx"
    arguments = ["run", "--profile", str(profile), "--db", str(database), "-o", str(report), "--pause", "0"]
    assert main(arguments) == 0
    assert "первый прогон" in capsys.readouterr().out
    shop.update(demo.changed())
    assert main(arguments) == 0
    printed = capsys.readouterr().out
    assert "Изменений:" in printed
    assert "подешевел" in printed
    assert load_workbook(report)["Прогоны"].max_row == 3


def test_profiles_command(capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["profiles"]) == 0
    printed = capsys.readouterr().out
    assert "books" in printed
    assert '"card"' in printed


def test_history_of_empty_database(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["history", "--db", str(tmp_path / "пусто.sqlite")]) == 0
    assert "нет прогонов" in capsys.readouterr().out


def test_unknown_profile_is_explained(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["run", "--profile", "ozon", "--db", str(tmp_path / "x.sqlite")]) == 2
    assert "не найден" in capsys.readouterr().err
